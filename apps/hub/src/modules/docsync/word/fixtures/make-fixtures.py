"""Sinh file Word mẫu cho test vàng của filter.lua.

Chạy lại khi đổi nội dung mẫu, rồi đổi `.odt` và duyệt lại `*.md` vàng bằng mắt:

    python make-fixtures.py
    soffice --headless --convert-to odt --outdir . srs.docx

Cần python-docx 1.2 (có `add_comment`). File sinh ra được commit, nên test không cần Python.
"""

import io
import struct
import zipfile
from copy import deepcopy
import zlib
from datetime import datetime
from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.shape import CT_Inline
from docx.shared import Emu, Pt, RGBColor

HERE = Path(__file__).resolve().parent
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "v": "urn:schemas-microsoft-com:vml",
}
DATE = "2026-10-19T00:00:00Z"
# Thời điểm cố định trong zip để chạy lại script ra cùng byte.
ZIP_TIME = (2026, 10, 19, 0, 0, 0)


def xml(body: str):
    decls = " ".join(f'xmlns:{p}="{u}"' for p, u in NS.items())
    root = parse_xml(f"<w:root {decls}>{body}</w:root>")
    return list(root)


def png(rgb: tuple[int, int, int], size: int = 8) -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    row = b"\x00" + bytes(rgb) * size
    ihdr = struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(row * size, 9)) + chunk(b"IEND", b"")


def emf() -> bytes:
    # EMF tối thiểu: EMR_HEADER (108 byte) + EMR_EOF (20 byte).
    total = 128
    header = struct.pack(
        "<II4i4iIIIIHHIII2i2iIII2i",
        1, 108, 0, 0, 15, 15, 0, 0, 400, 400, 0x464D4520, 0x10000, total, 2, 1, 0, 0, 0, 0,
        1024, 768, 320, 240, 0, 0, 0, 320000, 240000,
    )
    eof = struct.pack("<IIIII", 14, 20, 0, 16, 20)
    return header + eof


def add_picture_run(paragraph, blob: bytes, width=Emu(762000)):
    run = paragraph.add_run()
    run.add_picture(io.BytesIO(blob), width=width)
    return run


def add_emf_run(document, paragraph, blob: bytes):
    # python-docx không nhận EMF ở add_picture, nên tự thêm phần ảnh và khung vẽ.
    part = Part(PackURI("/word/media/image10.emf"), "image/x-emf", blob, document.part.package)
    rid = document.part.relate_to(part, RT.IMAGE)
    inline = CT_Inline.new_pic_inline(900, rid, "image10.emf", Emu(762000), Emu(762000))
    run = paragraph.add_run()
    run._r.add_drawing(inline)


def bookmark(paragraph, name: str, bid: int):
    start, end = xml(f'<w:bookmarkStart w:id="{bid}" w:name="{name}"/><w:bookmarkEnd w:id="{bid}"/>')
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def internal_link(paragraph, anchor: str, text: str):
    (link,) = xml(f'<w:hyperlink w:anchor="{anchor}"><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:hyperlink>')
    paragraph._p.append(link)


def external_link(document, paragraph, url: str, text: str):
    rid = document.part.relate_to(url, RT.HYPERLINK, is_external=True)
    (link,) = xml(f'<w:hyperlink r:id="{rid}"><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:hyperlink>')
    paragraph._p.append(link)


def shade(run, fill: str):
    run._r.get_or_add_rPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))


def header_rows(table, count: int):
    for row in table.rows[:count]:
        row._tr.get_or_add_trPr().append(parse_xml(f"<w:tblHeader {nsdecls('w')}/>"))


def fill_cell(cell, text: str):
    cell.paragraphs[0].text = text


def new_numbering(document, style_name: str) -> int:
    # Cột STT là một danh sách riêng bắt đầu từ 1, như khi BA chèn bảng mới trong Word. Dùng định nghĩa số riêng thay
    # cho startOverride, vì pandoc áp startOverride cho từng ô và mọi ô ra số 1.
    num_id = document.styles[style_name].element.pPr.numPr.numId.val
    numbering = document.part.numbering_part.element
    source_id = numbering.num_having_numId(num_id).abstractNumId.val
    abstracts = numbering.findall(qn("w:abstractNum"))
    source = next(a for a in abstracts if a.get(qn("w:abstractNumId")) == str(source_id))
    copy = deepcopy(source)
    new_id = max(int(a.get(qn("w:abstractNumId"))) for a in abstracts) + 1
    copy.set(qn("w:abstractNumId"), str(new_id))
    for child in copy.findall(qn("w:nsid")) + copy.findall(qn("w:tmpl")):
        copy.remove(child)
    abstracts[-1].addnext(copy)
    return numbering.add_num(new_id).numId


def set_numbering(paragraph, num_id: int):
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


def heading_style(document, name: str, base: str):
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles[base]
    style.font.size = Pt(13)
    return style


def save(document, name: str):
    buffer = io.BytesIO()
    document.save(buffer)
    # Ghi lại zip với thời điểm cố định để file mẫu không đổi khi chạy lại script.
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source, zipfile.ZipFile(
        HERE / name, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, ZIP_TIME)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(fixed, source.read(info.filename))


def fixed_core(document):
    props = document.core_properties
    props.author = "xDev Hive"
    props.last_modified_by = "xDev Hive"
    props.revision = 1
    props.created = datetime(2026, 10, 19)
    props.modified = datetime(2026, 10, 19)


def build_srs():
    d = docx.Document()
    fixed_core(d)
    h1 = heading_style(d, "FIS_Heading1", "Heading 1")
    h3 = heading_style(d, "1b", "Heading 3")

    # Trang bìa, bảng ghi nhận thay đổi và mục lục như mẫu SRS của FIS.
    add_picture_run(d.add_paragraph(), png((200, 30, 30)))
    d.add_paragraph("TÀI LIỆU ĐẶC TẢ YÊU CẦU PHẦN MỀM", style="Title")
    d.add_paragraph("Mã tài liệu: eHosV2i_SRS_OPD.OPD99")
    info = d.add_table(rows=2, cols=2)
    info.style = "Table Grid"
    for r, (k, v) in enumerate([("Dự án", "eHospital V2i"), ("Phiên bản", "1.1")]):
        fill_cell(info.cell(r, 0), k)
        fill_cell(info.cell(r, 1), v)
    d.add_paragraph("BẢNG GHI NHẬN THAY ĐỔI")
    log = d.add_table(rows=2, cols=3)
    log.style = "Table Grid"
    for c, text in enumerate(["Ngày", "Người sửa", "Nội dung"]):
        fill_cell(log.cell(0, c), text)
    for c, text in enumerate(["01/10/2026", "BA", "Tạo mới"]):
        fill_cell(log.cell(1, c), text)
    d.add_paragraph("MỤC LỤC")
    for anchor, text in [("_Toc100001", "I. TỔNG QUAN\t3"), ("_Toc100002", "II. QUY TRÌNH NGHIỆP VỤ\t4"),
                         ("_Toc100003", "III. YÊU CẦU\t5")]:
        internal_link(d.add_paragraph(), anchor, text)

    p = d.add_paragraph("I. TỔNG QUAN", style=h1)
    bookmark(p, "_Toc100001", 1)
    d.add_paragraph("1. Mục đích", style="Heading 2")
    p = d.add_paragraph("Đoạn có ")
    p.add_run("chữ đậm").bold = True
    p.add_run(", ")
    p.add_run("chữ nghiêng").italic = True
    p.add_run(", ")
    p.add_run("gạch chân").underline = True
    p.add_run(", ")
    p.add_run("chữ đỏ").font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    p.add_run(", ")
    p.add_run("tô vàng").font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.add_run(", ")
    shade(p.add_run("nền xám"), "D9D9D9")
    p.add_run(", link ngoài ")
    external_link(d, p, "https://example.com/srs", "trang SRS")
    p.add_run(" và link nội bộ ")
    internal_link(p, "_Ref200001", "mục 1.1")
    p.add_run(".")

    p = d.add_paragraph("1.1. Phạm vi", style=h3)
    bookmark(p, "_Ref200001", 2)
    d.add_paragraph("Danh sách đánh số có mục rỗng:")
    d.add_paragraph("Bước một", style="List Number")
    d.add_paragraph("", style="List Number")
    d.add_paragraph("Bước hai", style="List Number")
    d.add_paragraph("   ", style="List Number")
    d.add_paragraph("Danh sách chấm có mục rỗng:")
    d.add_paragraph("Ý một", style="List Bullet")
    d.add_paragraph("", style="List Bullet")
    d.add_paragraph("Ý hai", style="List Bullet")
    d.add_paragraph("")
    d.add_paragraph("   ")

    # Theo dõi thay đổi: chấp nhận cả phần thêm và phần xoá.
    p = d.add_paragraph("Câu có thay đổi: ")
    for el in xml(
        f'<w:ins w:id="10" w:author="BA" w:date="{DATE}"><w:r><w:t xml:space="preserve">chữ được thêm </w:t></w:r></w:ins>'
        f'<w:del w:id="11" w:author="BA" w:date="{DATE}"><w:r><w:delText xml:space="preserve">chữ bị xoá </w:delText></w:r></w:del>'
        '<w:r><w:t>còn lại.</w:t></w:r>'
    ):
        p._p.append(el)
    p = d.add_paragraph("Câu có ")
    commented = p.add_run("bình luận")
    p.add_run(" của người duyệt.")
    comment = d.add_comment(commented, text="Cần làm rõ.", author="Reviewer", initials="RV")
    comment._comment_elm.set(qn("w:date"), DATE)

    p = d.add_paragraph("II. QUY TRÌNH NGHIỆP VỤ", style=h1)
    bookmark(p, "_Toc100002", 3)
    add_picture_run(d.add_paragraph(), png((30, 90, 200)))
    d.add_paragraph("Hình 1: Quy trình khám bệnh", style="Caption")
    d.add_paragraph("Sơ đồ dạng EMF:")
    add_emf_run(d, d.add_paragraph(), emf())

    p = d.add_paragraph("III. YÊU CẦU", style=h1)
    bookmark(p, "_Toc100003", 4)
    d.add_paragraph("Bảng 1: Danh sách trường", style="Caption")
    t = d.add_table(rows=6, cols=4)
    t.style = "Table Grid"
    header_rows(t, 1)
    for c, text in enumerate(["STT", "Trường", "Mô tả", "Bắt buộc"]):
        fill_cell(t.cell(0, c), text)
    # Cột STT đánh số tự động liền qua các ô.
    stt = new_numbering(d, "List Number")
    for r in range(1, 6):
        t.cell(r, 0).paragraphs[0].style = d.styles["List Number"]
        set_numbering(t.cell(r, 0).paragraphs[0], stt)
    fill_cell(t.cell(1, 1), "Mã bệnh nhân")
    fill_cell(t.cell(1, 2), "Sinh tự động | không sửa")
    fill_cell(t.cell(1, 3), "Có")
    # Gộp ngang: mô tả và cột bắt buộc.
    merged = t.cell(2, 2).merge(t.cell(2, 3))
    fill_cell(t.cell(2, 1), "Họ tên")
    fill_cell(merged, "Ô gộp ngang")
    # Gộp dọc: trường dùng chung cho hai hàng.
    vmerged = t.cell(3, 1).merge(t.cell(4, 1))
    fill_cell(vmerged, "Địa chỉ")
    cell = t.cell(3, 2)
    fill_cell(cell, "Đoạn một")
    cell.add_paragraph("Đoạn hai")
    run = cell.add_paragraph("Dòng trên").add_run()
    run.add_break()
    run.add_text("dòng dưới")
    fill_cell(t.cell(3, 3), "Có")
    cell = t.cell(4, 2)
    fill_cell(cell, "Gồm:")
    cell.add_paragraph("Tỉnh", style="List Bullet")
    cell.add_paragraph("", style="List Bullet")
    cell.add_paragraph("Xã", style="List Bullet")
    fill_cell(t.cell(4, 3), "Không")
    fill_cell(t.cell(5, 1), "Ảnh")
    add_picture_run(t.cell(5, 2).paragraphs[0], png((30, 160, 60)), width=Emu(381000))
    cell = t.cell(5, 3)
    fill_cell(cell, "Bảng con:")
    inner = cell.add_table(rows=2, cols=2)
    inner.style = "Table Grid"
    for (r, c), text in {(0, 0): "Mã", (0, 1): "Tên", (1, 0): "01", (1, 1): "Hà Nội"}.items():
        fill_cell(inner.cell(r, c), text)
    cell.add_paragraph("Sau bảng con")

    d.add_paragraph("Bảng có hai hàng tiêu đề, không chú thích:")
    t = d.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    header_rows(t, 2)
    for (r, c), text in {(0, 0): "Nhóm", (0, 1): "Chi tiết", (1, 0): "Mã", (1, 1): "Tên", (2, 0): "A", (2, 1): "Khám"}.items():
        fill_cell(t.cell(r, c), text)
    d.add_paragraph("Hết tài liệu.")
    save(d, "srs.docx")


def script_run(paragraph, text: str, kind: str, bold: bool = False):
    # kind: "subscript" hoặc "superscript"; python-docx đặt w:vertAlign qua thuộc tính của run.
    run = paragraph.add_run(text)
    setattr(run.font, kind, True)
    run.bold = bold or None
    return run


def build_inline_lists():
    d = docx.Document()
    fixed_core(d)
    d.add_paragraph("I. TỔNG QUAN", style="Heading 1")

    p = d.add_paragraph("Chỉ số SpO")
    script_run(p, "2", "subscript")
    p.add_run(" và diện tích m")
    script_run(p, "2", "superscript")
    p.add_run(" ngoài bảng.")

    d.add_paragraph("Bảng chỉ số:")
    t = d.add_table(rows=7, cols=2)
    t.style = "Table Grid"
    header_rows(t, 1)
    fill_cell(t.cell(0, 0), "Chỉ số")
    fill_cell(t.cell(0, 1), "Đơn vị")
    p = t.cell(1, 0).paragraphs[0]
    p.add_run("SpO")
    script_run(p, "2", "subscript")
    fill_cell(t.cell(1, 1), "%")
    fill_cell(t.cell(2, 0), "Diện tích")
    p = t.cell(2, 1).paragraphs[0]
    p.add_run("m")
    script_run(p, "2", "superscript")
    # Chỉ số có dấu và ngoặc, trong ô nhiều đoạn.
    p = t.cell(3, 0).paragraphs[0]
    p.add_run("Ion Ca")
    script_run(p, "2+", "superscript")
    p = t.cell(3, 1).paragraphs[0]
    p.add_run("Dòng một")
    p = t.cell(3, 1).add_paragraph("H")
    script_run(p, "2", "subscript")
    p.add_run("O")
    # Chỉ số in đậm.
    p = t.cell(4, 0).paragraphs[0]
    p.add_run("x")
    script_run(p, "(n+1)", "superscript", bold=True)
    fill_cell(t.cell(4, 1), "Mũ")
    # Chỉ số có chữ không đổi được sang Unicode: giữ chữ, bỏ định dạng.
    p = t.cell(5, 0).paragraphs[0]
    p.add_run("T")
    script_run(p, "max", "subscript")
    p = t.cell(5, 1).paragraphs[0]
    p.add_run("Lần 5")
    script_run(p, "th", "superscript")
    # Chỉ số trong tiêu đề cột.
    p = t.cell(6, 0).paragraphs[0]
    p.add_run("Nồng độ CO")
    script_run(p, "2", "subscript")
    fill_cell(t.cell(6, 1), "ppm")
    # Tiêu đề cột thứ hai có chỉ số: đặt sau để hàng đầu giữ chữ thường cho cả hai bản.
    p = t.cell(0, 1).paragraphs[0]
    script_run(p, "(2)", "superscript")

    # Hai danh sách đánh số liền nhau, mỗi danh sách một định nghĩa số riêng nên Word đánh lại từ 1.
    d.add_paragraph("Hai danh sách đánh số liền nhau:")
    for numbering, items in [(new_numbering(d, "List Number"), ["Bước một", "Bước hai"]),
                             (new_numbering(d, "List Number"), ["Bước ba", "Bước bốn"])]:
        for text in items:
            q = d.add_paragraph(text, style="List Number")
            set_numbering(q, numbering)
    d.add_paragraph("Hai danh sách chấm liền nhau:")
    for numbering, items in [(new_numbering(d, "List Bullet"), ["Ý một", "Ý hai"]),
                             (new_numbering(d, "List Bullet"), ["Ý ba"])]:
        for text in items:
            q = d.add_paragraph(text, style="List Bullet")
            set_numbering(q, numbering)
    d.add_paragraph("Hết tài liệu.")
    save(d, "inline-lists.docx")


FOOTNOTES = (
    '<w:footnotes {decls}>'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    '<w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> Ghi chú cuối trang.</w:t></w:r></w:p></w:footnote>'
    "</w:footnotes>"
)


def build_no_h1():
    d = docx.Document()
    fixed_core(d)
    d.add_paragraph("Biên bản không có tiêu đề cấp 1", style="Heading 2")
    p = d.add_paragraph("Câu có chú thích cuối trang")
    decls = " ".join(f'xmlns:{p}="{u}"' for p, u in NS.items())
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        FOOTNOTES.format(decls=decls).encode("utf-8"),
        d.part.package,
    )
    d.part.relate_to(part, RT.FOOTNOTES)
    (ref,) = xml('<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteReference w:id="1"/></w:r>')
    p._p.append(ref)
    p.add_run(".")
    p = d.add_paragraph("Công thức: ")
    (math,) = xml("<m:oMath><m:r><m:t>x+1=2</m:t></m:r></m:oMath>")
    p._p.append(math)
    p = d.add_paragraph("Hộp chữ: ")
    (box,) = xml(
        '<w:r><w:pict><v:shape style="width:100pt;height:30pt"><v:textbox><w:txbxContent>'
        "<w:p><w:r><w:t>Chữ trong hộp</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
    )
    p._p.append(box)
    d.add_paragraph("Kết thúc.")
    save(d, "no-h1.docx")


BUILDERS = {"srs": build_srs, "no-h1": build_no_h1, "inline-lists": build_inline_lists}

if __name__ == "__main__":
    # Chạy riêng một mẫu được, vì python-docx/lxml khác phiên bản có thể ra byte khác cho các mẫu còn lại.
    import sys

    for name in sys.argv[1:] or BUILDERS:
        BUILDERS[name]()
